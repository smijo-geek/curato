"""Extraction port — turns a source file into (records, raw_text).

A source can be anything: PDF, Word, Excel, CSV/TSV/delimited, text, XML, images, video, COBOL
copybooks, or a system like SAP or Salesforce. Each source type is one `ExtractionPort`; the
`ExtractionRegistry` dispatches a file to the first extractor that `supports()` it. Adding a new
source is therefore a single new class — the agents never change.

Shipped here: a rules-based PDF invoice reader and a generic CSV/TSV reader (both stdlib / pypdf,
offline). See docs/adapters.md for the pattern; the cloud/LLM extractors (Docling, Azure Doc
Intelligence, Textract, Unstructured) implement the same two methods.
"""
from __future__ import annotations
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from typing import List, Dict, Any
import re


@dataclass
class Extraction:
    records: List[Dict[str, Any]] = field(default_factory=list)
    raw_text: str = ""


class ExtractionPort(ABC):
    @abstractmethod
    def supports(self, filename: str) -> bool: ...
    @abstractmethod
    def extract(self, filename: str, data: bytes) -> Extraction: ...


class PdfInvoiceExtractor(ExtractionPort):
    """Rules-based invoice reader. In production the LLM/vision model does this step, but the
    rules make the demo deterministic and offline."""

    FIELDS = {
        "invoice_number": re.compile(r"Invoice Number:\s*(\S+)", re.I),
        "invoice_date":   re.compile(r"Date:\s*(\d{4}-\d{2}-\d{2})", re.I),
        "vendor":         re.compile(r"Vendor:\s*(.+)", re.I),
        "_amount":        re.compile(r"Amount Due:\s*([\d,]+\.\d{2})\s*([A-Z]{3})", re.I),
    }

    def supports(self, filename: str) -> bool:
        return filename.lower().endswith(".pdf")

    def _read_text(self, data: bytes) -> str:
        from io import BytesIO
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    def extract(self, filename: str, data: bytes) -> Extraction:
        text = self._read_text(data)
        rec: Dict[str, Any] = {}
        for key, rx in self.FIELDS.items():
            m = rx.search(text)
            if not m:
                continue
            if key == "_amount":
                rec["amount"] = float(m.group(1).replace(",", ""))
                rec["currency"] = m.group(2).upper()
            else:
                rec[key] = m.group(1).strip()
        rec["source_file"] = filename
        return Extraction(records=[rec] if rec.get("invoice_number") else [], raw_text=text)


def _coerce(v):
    """Best-effort type inference for a cell: int, then float, else the original string."""
    if v is None or v == "":
        return None
    try:
        return int(v)
    except (ValueError, TypeError):
        pass
    try:
        return float(v)
    except (ValueError, TypeError):
        return v


class CsvExtractor(ExtractionPort):
    """Generic tabular reader for CSV / TSV / delimited text. Each row becomes a record; column
    names flow straight into the Profiler and the glossary Mapping. Numeric cells are coerced so
    downstream aggregations get real types."""

    def supports(self, filename: str) -> bool:
        return filename.lower().endswith((".csv", ".tsv"))

    def extract(self, filename: str, data: bytes) -> Extraction:
        import csv
        from io import StringIO
        text = data.decode("utf-8-sig", errors="replace")
        delim = "\t" if filename.lower().endswith(".tsv") else ","
        rows = list(csv.DictReader(StringIO(text), delimiter=delim))
        for r in rows:
            for k in list(r.keys()):
                r[k] = _coerce(r[k])
            r["source_file"] = filename
        return Extraction(records=rows, raw_text=text)


class ExcelExtractor(ExtractionPort):
    """Reads .xlsx / .xlsm: the first row is the header, each later row a record. openpyxl returns
    typed cells, so no coercion is needed. Needs the `xlsx` extra (openpyxl), imported lazily so the
    package still works without it."""

    def supports(self, filename: str) -> bool:
        return filename.lower().endswith((".xlsx", ".xlsm"))

    def extract(self, filename: str, data: bytes) -> Extraction:
        from io import BytesIO
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise RuntimeError("Excel source needs openpyxl — install with: pip install 'curato[xlsx]'") from e
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(h) if h is not None else f"col{i}" for i, h in enumerate(next(it))]
        except StopIteration:
            return Extraction(records=[], raw_text="")
        records: List[Dict[str, Any]] = []
        for row in it:
            if row is None or all(c is None for c in row):
                continue
            rec = {header[i]: row[i] for i in range(min(len(header), len(row)))}
            rec["source_file"] = filename
            records.append(rec)
        wb.close()
        text = "\n".join(" ".join("" if v is None else str(v) for v in r.values()) for r in records)
        return Extraction(records=records, raw_text=text)


class ExtractionRegistry(ExtractionPort):
    """Holds many extractors and routes each file to the first that supports it. This is the
    'source can be anything' seam: register more extractors, nothing else changes."""

    def __init__(self, extractors: List[ExtractionPort]):
        self.extractors = extractors

    def _for(self, filename: str) -> ExtractionPort | None:
        return next((e for e in self.extractors if e.supports(filename)), None)

    def supports(self, filename: str) -> bool:
        return self._for(filename) is not None

    def extract(self, filename: str, data: bytes) -> Extraction:
        ex = self._for(filename)
        if ex is None:
            return Extraction(records=[], raw_text="")
        return ex.extract(filename, data)
